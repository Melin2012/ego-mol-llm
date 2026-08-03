#!/usr/bin/env python3
"""
Export full per-spectrum results table for MSG HNSW test64 Grok handout.

Outcome categories (mutually exclusive, highest wins):
  1. exact              — canonical SMILES equal
  2. same_connectivity  — IK1 equal, SMILES not exact (stereo / tautomer / canon form)
  3. same_formula_isomer— same molecular formula, different IK1 (isomer / regio / stereo-class)
  4. close_T085         — Morgan Tanimoto ≥ 0.85 (not above)
  5. close_T070         — 0.70 ≤ T < 0.85
  6. moderate_T050      — 0.50 ≤ T < 0.70
  7. weak_T030          — 0.30 ≤ T < 0.50
  8. distant_miss       — has pred SMILES, T < 0.30 (or unparseable T)
  9. empty              — empty SMILES
 10. missing_pred       — no prediction file
 11. invalid_smiles     — non-empty SMILES that RDKit cannot parse

Also exports flags: mass_ok, formula_match, ik1, exact, T bands, etc.
"""

from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rdkit import Chem, RDLogger
from rdkit.Chem import AllChem, DataStructs, Descriptors, rdMolDescriptors

RDLogger.DisableLog("rdApp.*")

HAND = Path(r"C:\Users\AlexeyMelnik\Desktop\MSG_HNSW_test64_nist_neigh_GROK_handout")
SEALED = Path(r"C:\Users\AlexeyMelnik\Desktop\MSG_HNSW_test64_nist_neigh_SEALED_truth")
PRED_DIR = HAND / "predictions"

OUT_XLSX = Path(
    r"C:\Users\AlexeyMelnik\Desktop\MSG_HNSW_test64_GROK_RESULTS_FULL.xlsx"
)
OUT_CSV = SEALED / "GROK_HANDOUT_TEST64_FULL_TABLE.csv"
OUT_XLSX_SEALED = SEALED / "GROK_HANDOUT_TEST64_FULL_TABLE.xlsx"
OUT_SUMMARY_CSV = SEALED / "GROK_HANDOUT_TEST64_CATEGORY_SUMMARY.csv"


def mol(s: str | None):
    if not s or not str(s).strip():
        return None
    return Chem.MolFromSmiles(str(s).strip())


def can(s: str | None) -> str | None:
    m = mol(s)
    return Chem.MolToSmiles(m) if m else None


def ik1(s: str | None) -> str | None:
    m = mol(s)
    if not m:
        return None
    return Chem.MolToInchiKey(m).split("-")[0].upper()


def full_ik(s: str | None) -> str | None:
    m = mol(s)
    return Chem.MolToInchiKey(m) if m else None


def formula(s: str | None) -> str | None:
    m = mol(s)
    return rdMolDescriptors.CalcMolFormula(m) if m else None


def tani(a: str | None, b: str | None) -> float | None:
    ma, mb = mol(a), mol(b)
    if not ma or not mb:
        return None
    fa = AllChem.GetMorganFingerprintAsBitVect(ma, 2, nBits=2048)
    fb = AllChem.GetMorganFingerprintAsBitVect(mb, 2, nBits=2048)
    return float(DataStructs.TanimotoSimilarity(fa, fb))


def exact_mass(s: str | None) -> float | None:
    m = mol(s)
    return float(Descriptors.ExactMolWt(m)) if m else None


def mass_fit(smiles: str | None, mz: float, tol: float = 0.05):
    m = mol(smiles)
    if not m:
        return False, None, None, None
    M = Descriptors.ExactMolWt(m)
    adducts = {
        "[M+H]+": M + 1.007825032,
        "[M-H]-": M - 1.007825032,
        "[M+Na]+": M + 22.98922127,
        "[M+K]+": M + 38.9637069,
        "[M+NH4]+": M + 18.03382555,
        "[M+H-H2O]+": M + 1.007825032 - 18.0105647,
        "[2M+H]+": 2 * M + 1.007825032,
        "[2M-H]-": 2 * M - 1.007825032,
        "[2M+Na]+": 2 * M + 22.98922127,
        "[2M+NH4]+": 2 * M + 18.03382555,
        "[2M+H-H2O]+": 2 * M + 1.007825032 - 18.0105647,
        "[3M+H]+": 3 * M + 1.007825032,
        "[M]+": M,
    }
    best_a, best_e, best_t = None, 1e9, None
    for a, tmz in adducts.items():
        e = abs(tmz - mz)
        if e < best_e:
            best_a, best_e, best_t = a, e, tmz
    return best_e <= tol, best_a, round(best_e, 5), None if best_t is None else round(best_t, 5)


def classify(
    *,
    missing: bool,
    empty: bool,
    invalid: bool,
    exact_ok: bool,
    ik_ok: bool,
    form_ok: bool,
    tt: float | None,
) -> str:
    if missing:
        return "missing_pred"
    if empty:
        return "empty"
    if invalid:
        return "invalid_smiles"
    if exact_ok:
        return "exact"
    if ik_ok:
        return "same_connectivity"  # stereo / tautomer / canon
    if form_ok:
        return "same_formula_isomer"
    if tt is not None and tt >= 0.85:
        return "close_T085"
    if tt is not None and tt >= 0.70:
        return "close_T070"
    if tt is not None and tt >= 0.50:
        return "moderate_T050"
    if tt is not None and tt >= 0.30:
        return "weak_T030"
    return "distant_miss"


CATEGORY_ORDER = [
    "exact",
    "same_connectivity",
    "same_formula_isomer",
    "close_T085",
    "close_T070",
    "moderate_T050",
    "weak_T030",
    "distant_miss",
    "empty",
    "invalid_smiles",
    "missing_pred",
]

CATEGORY_HELP = {
    "exact": "Canonical SMILES identical to sealed truth",
    "same_connectivity": "InChIKey first block (IK1) match; SMILES not exact (stereo/tautomer/representation)",
    "same_formula_isomer": "Same molecular formula, different connectivity (isomer / regioisomer)",
    "close_T085": "Morgan-2 Tanimoto ≥ 0.85; not same formula and not IK1",
    "close_T070": "0.70 ≤ Tanimoto < 0.85 — close analog",
    "moderate_T050": "0.50 ≤ Tanimoto < 0.70 — moderate similarity",
    "weak_T030": "0.30 ≤ Tanimoto < 0.50 — weak similarity",
    "distant_miss": "Parsed prediction with T < 0.30 (or no T) — clear structural miss",
    "empty": "Empty predicted SMILES",
    "invalid_smiles": "Non-empty SMILES RDKit cannot parse",
    "missing_pred": "No prediction JSON file",
}

# Excel colors by category
CAT_FILL = {
    "exact": "C6EFCE",
    "same_connectivity": "A9D08E",
    "same_formula_isomer": "FFE699",
    "close_T085": "F8CBAD",
    "close_T070": "F4B183",
    "moderate_T050": "FCE4D6",
    "weak_T030": "DDEBF7",
    "distant_miss": "FFC7CE",
    "empty": "D9D9D9",
    "invalid_smiles": "C00000",
    "missing_pred": "808080",
}


def main() -> int:
    truth_rows = list(
        csv.DictReader((SEALED / "truth_index.csv").open(encoding="utf-8-sig", newline=""))
    )
    truth = {r["spectrum_id"]: r for r in truth_rows}

    rows: list[dict] = []
    for sid in sorted(truth):
        t = truth[sid]
        ppath = PRED_DIR / f"{sid}.json"
        if not ppath.exists():
            pred = {}
            missing = True
        else:
            pred = json.loads(ppath.read_text(encoding="utf-8-sig"))
            missing = False

        ps = (pred.get("smiles") or "").strip()
        ts = (t.get("true_smiles") or "").strip()
        empty = not bool(ps) and not missing
        pm = mol(ps) if ps else None
        invalid = bool(ps) and pm is None

        tik_sealed = (t.get("true_inchikey") or "").split("-")[0].upper()
        tik = ik1(ts) or tik_sealed or ""
        pik = ik1(ps) or ""
        pcs = can(ps)
        tcs = can(ts)
        exact_ok = bool(pcs and tcs and pcs == tcs)
        ik_ok = bool(pik and tik and pik == tik)
        pf = formula(ps)
        tf = formula(ts) or (t.get("true_formula") or "").strip() or None
        form_ok = bool(pf and tf and pf == tf)
        tt = tani(ps, ts) if pm and mol(ts) else None

        try:
            mz = float(t.get("precursor_mz") or 0)
        except Exception:
            mz = 0.0
        p_mok, p_best_a, p_res, p_theo = mass_fit(ps, mz) if ps and not invalid else (False, None, None, None)
        t_mok, t_best_a, t_res, t_theo = mass_fit(ts, mz) if ts else (False, None, None, None)

        tmw = exact_mass(ts)
        pmw = exact_mass(ps) if not invalid else None
        dmw = None if tmw is None or pmw is None else round(pmw - tmw, 4)

        cat = classify(
            missing=missing,
            empty=empty,
            invalid=invalid,
            exact_ok=exact_ok,
            ik_ok=ik_ok,
            form_ok=form_ok,
            tt=tt,
        )

        # secondary labels for filtering
        is_hit = cat in ("exact", "same_connectivity")
        is_near = cat in (
            "exact",
            "same_connectivity",
            "same_formula_isomer",
            "close_T085",
            "close_T070",
        )
        is_miss = cat in ("distant_miss", "empty", "invalid_smiles", "missing_pred")
        is_partial = cat in (
            "same_formula_isomer",
            "close_T085",
            "close_T070",
            "moderate_T050",
            "weak_T030",
        )

        rat = pred.get("rationale") or ""
        try:
            conf = float(pred.get("confidence"))
        except (TypeError, ValueError):
            conf = None

        rows.append(
            {
                "spectrum_id": sid,
                "hnsw_index": t.get("hnsw_index"),
                "msg_identifier": t.get("msg_identifier"),
                "fold": t.get("fold"),
                "precursor_mz": t.get("precursor_mz"),
                "true_adduct_meta": t.get("adduct"),
                "outcome_category": cat,
                "outcome_rank": CATEGORY_ORDER.index(cat) if cat in CATEGORY_ORDER else 99,
                "is_structure_hit_IK1_or_exact": is_hit,
                "is_near_miss_or_better": is_near,
                "is_partial_credit": is_partial,
                "is_clear_miss": is_miss,
                "exact": exact_ok,
                "ik1_match": ik_ok,
                "formula_match": form_ok,
                "tanimoto_morgan2": None if tt is None else round(tt, 4),
                "tani_ge_0.85": bool(tt is not None and tt >= 0.85),
                "tani_ge_0.70": bool(tt is not None and tt >= 0.70),
                "tani_ge_0.50": bool(tt is not None and tt >= 0.50),
                "tani_ge_0.30": bool(tt is not None and tt >= 0.30),
                "true_formula": tf,
                "pred_formula": pf or pred.get("formula"),
                "true_ik1": tik,
                "pred_ik1": pik,
                "true_inchikey_full": full_ik(ts) or "",
                "pred_inchikey_full": full_ik(ps) or "",
                "true_exact_mass": None if tmw is None else round(tmw, 5),
                "pred_exact_mass": None if pmw is None else round(pmw, 5),
                "delta_exact_mass_pred_minus_true": dmw,
                "true_mass_ok_vs_precursor": t_mok,
                "true_best_adduct": t_best_a,
                "true_mass_resid_da": t_res,
                "pred_mass_ok_vs_precursor": p_mok,
                "pred_best_adduct_fit": p_best_a,
                "pred_mass_resid_da": p_res,
                "pred_stated_adduct": pred.get("adduct"),
                "confidence": conf,
                "pred_name": pred.get("iupac_or_common_name"),
                "true_smiles": ts,
                "pred_smiles": ps,
                "true_smiles_canonical": tcs or "",
                "pred_smiles_canonical": pcs or "",
                "rationale_len": len(rat),
                "rationale": rat,
                "model": pred.get("model"),
                "method": pred.get("method"),
                "n_neighbor_nist_hits": t.get("n_neighbor_nist_hits"),
                "instrument_type": t.get("instrument_type"),
                "collision_energy": t.get("collision_energy"),
                "missing_pred": missing,
                "empty_smiles": empty,
                "invalid_smiles": invalid,
            }
        )

    # sort: misses first for review? User wants to see misses — provide two sheets
    # default sheet: by category order then spectrum_id
    rows_sorted = sorted(rows, key=lambda r: (r["outcome_rank"], r["spectrum_id"]))
    misses = [r for r in rows_sorted if r["outcome_category"] not in ("exact", "same_connectivity")]
    hits = [r for r in rows_sorted if r["outcome_category"] in ("exact", "same_connectivity")]
    near = [
        r
        for r in rows_sorted
        if r["outcome_category"]
        in ("same_formula_isomer", "close_T085", "close_T070", "moderate_T050")
    ]

    counts = Counter(r["outcome_category"] for r in rows)
    n = len(rows)

    # --- CSV full ---
    fieldnames = list(rows[0].keys())
    with OUT_CSV.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows_sorted)

    with OUT_SUMMARY_CSV.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(
            f, fieldnames=["outcome_category", "n", "pct", "description"]
        )
        w.writeheader()
        for cat in CATEGORY_ORDER:
            c = counts.get(cat, 0)
            if c or cat in counts:
                w.writerow(
                    {
                        "outcome_category": cat,
                        "n": c,
                        "pct": round(100 * c / n, 1) if n else 0,
                        "description": CATEGORY_HELP.get(cat, ""),
                    }
                )
        # also aggregate rows
        w.writerow({})
        w.writerow(
            {
                "outcome_category": "AGG_structure_hit_exact_or_IK1",
                "n": sum(1 for r in rows if r["is_structure_hit_IK1_or_exact"]),
                "pct": round(
                    100
                    * sum(1 for r in rows if r["is_structure_hit_IK1_or_exact"])
                    / n,
                    1,
                ),
                "description": "exact + same_connectivity",
            }
        )
        w.writerow(
            {
                "outcome_category": "AGG_near_or_better_T070_or_isomer_or_hit",
                "n": sum(1 for r in rows if r["is_near_miss_or_better"]),
                "pct": round(
                    100 * sum(1 for r in rows if r["is_near_miss_or_better"]) / n, 1
                ),
                "description": "hit + same_formula_isomer + close_T085 + close_T070",
            }
        )
        w.writerow(
            {
                "outcome_category": "AGG_formula_match_any",
                "n": sum(1 for r in rows if r["formula_match"]),
                "pct": round(100 * sum(1 for r in rows if r["formula_match"]) / n, 1),
                "description": "any formula match (includes exact/IK1/isomer)",
            }
        )
        w.writerow(
            {
                "outcome_category": "AGG_pred_mass_ok",
                "n": sum(1 for r in rows if r["pred_mass_ok_vs_precursor"]),
                "pct": round(
                    100 * sum(1 for r in rows if r["pred_mass_ok_vs_precursor"]) / n, 1
                ),
                "description": "prediction fits precursor under common adducts (0.05 Da)",
            }
        )

    # --- Excel ---
    def style_header(ws):
        fill = PatternFill("solid", fgColor="1F4E79")
        font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"

    def write_sheet(ws, data: list[dict], columns: list[str] | None = None):
        cols = columns or fieldnames
        ws.append(cols)
        for r in data:
            ws.append([r.get(c, "") for c in cols])
        style_header(ws)
        # color outcome column if present
        if "outcome_category" in cols:
            ci = cols.index("outcome_category") + 1
            for ri, r in enumerate(data, start=2):
                cat = r.get("outcome_category")
                color = CAT_FILL.get(cat)
                if color:
                    ws.cell(ri, ci).fill = PatternFill("solid", fgColor=color)
        # widths
        for i, c in enumerate(cols, 1):
            width = 14
            if "smiles" in c:
                width = 40
            elif c == "rationale":
                width = 50
            elif "name" in c:
                width = 28
            elif c == "outcome_category":
                width = 20
            ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
        thin = Border(
            left=Side(style="thin", color="DDDDDD"),
            right=Side(style="thin", color="DDDDDD"),
            top=Side(style="thin", color="DDDDDD"),
            bottom=Side(style="thin", color="DDDDDD"),
        )
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(cols)):
            for cell in row:
                cell.border = thin
                cell.alignment = Alignment(vertical="top", wrap_text=False)

    # compact columns for readable main view
    compact_cols = [
        "spectrum_id",
        "outcome_category",
        "is_structure_hit_IK1_or_exact",
        "is_partial_credit",
        "exact",
        "ik1_match",
        "formula_match",
        "tanimoto_morgan2",
        "confidence",
        "true_formula",
        "pred_formula",
        "true_ik1",
        "pred_ik1",
        "pred_name",
        "pred_mass_ok_vs_precursor",
        "pred_mass_resid_da",
        "pred_stated_adduct",
        "pred_best_adduct_fit",
        "precursor_mz",
        "delta_exact_mass_pred_minus_true",
        "true_smiles",
        "pred_smiles",
        "msg_identifier",
        "hnsw_index",
        "n_neighbor_nist_hits",
        "rationale_len",
    ]

    wb = Workbook()

    # Summary sheet
    ws0 = wb.active
    ws0.title = "Summary"
    ws0["A1"] = "MSG HNSW test64 — Grok free-form full results"
    ws0["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws0["A2"] = f"Handout: {HAND}"
    ws0["A3"] = f"Predictions: {PRED_DIR}"
    ws0["A4"] = f"n = {n}"
    ws0["A6"] = "outcome_category"
    ws0["B6"] = "n"
    ws0["C6"] = "pct"
    ws0["D6"] = "description"
    for cell in ws0[6]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    ri = 7
    for cat in CATEGORY_ORDER:
        c = counts.get(cat, 0)
        ws0.cell(ri, 1, cat)
        ws0.cell(ri, 2, c)
        ws0.cell(ri, 3, round(100 * c / n, 1) if n else 0)
        ws0.cell(ri, 4, CATEGORY_HELP.get(cat, ""))
        if cat in CAT_FILL:
            ws0.cell(ri, 1).fill = PatternFill("solid", fgColor=CAT_FILL[cat])
        ri += 1
    ri += 1
    ws0.cell(ri, 1, "Headline metrics").font = Font(bold=True)
    ri += 1
    metrics = [
        ("IK1 or exact (structure hit)", sum(1 for r in rows if r["is_structure_hit_IK1_or_exact"])),
        ("exact SMILES", sum(1 for r in rows if r["exact"])),
        ("same_connectivity only", counts.get("same_connectivity", 0)),
        ("same_formula_isomer", counts.get("same_formula_isomer", 0)),
        ("close analog T≥0.70 (close_T070+close_T085)", counts.get("close_T070", 0) + counts.get("close_T085", 0)),
        ("near-or-better (hit+isomer+T≥0.70)", sum(1 for r in rows if r["is_near_miss_or_better"])),
        ("any formula match", sum(1 for r in rows if r["formula_match"])),
        ("T ≥ 0.70 (any, incl hits)", sum(1 for r in rows if r["tani_ge_0.70"])),
        ("T ≥ 0.50 (any, incl hits)", sum(1 for r in rows if r["tani_ge_0.50"])),
        ("pred mass-OK", sum(1 for r in rows if r["pred_mass_ok_vs_precursor"])),
        ("distant_miss", counts.get("distant_miss", 0)),
    ]
    for name, val in metrics:
        ws0.cell(ri, 1, name)
        ws0.cell(ri, 2, val)
        ws0.cell(ri, 3, round(100 * val / n, 1) if n else 0)
        ri += 1
    ws0.column_dimensions["A"].width = 55
    ws0.column_dimensions["B"].width = 10
    ws0.column_dimensions["C"].width = 10
    ws0.column_dimensions["D"].width = 70

    ws1 = wb.create_sheet("All_results")
    write_sheet(ws1, rows_sorted, compact_cols)

    ws2 = wb.create_sheet("Misses_and_partial")
    write_sheet(ws2, misses, compact_cols)

    ws3 = wb.create_sheet("Hits_only")
    write_sheet(ws3, hits, compact_cols)

    ws4 = wb.create_sheet("Partial_near_misses")
    write_sheet(ws4, near, compact_cols)

    ws5 = wb.create_sheet("All_columns")
    write_sheet(ws5, rows_sorted, fieldnames)

    # legend
    wsL = wb.create_sheet("Category_legend")
    wsL.append(["outcome_category", "description", "color"])
    for cat in CATEGORY_ORDER:
        wsL.append([cat, CATEGORY_HELP[cat], CAT_FILL.get(cat, "")])
        if cat in CAT_FILL:
            wsL.cell(wsL.max_row, 1).fill = PatternFill("solid", fgColor=CAT_FILL[cat])
    style_header(wsL)
    wsL.column_dimensions["A"].width = 22
    wsL.column_dimensions["B"].width = 80

    for path in (OUT_XLSX, OUT_XLSX_SEALED):
        wb.save(path)

    print("=" * 64)
    print("TEST64 GROK — FULL RESULTS TABLE")
    print("=" * 64)
    print(f"n = {n}")
    for cat in CATEGORY_ORDER:
        c = counts.get(cat, 0)
        if c:
            print(f"  {cat:22s}  {c:3d}  ({100*c/n:5.1f}%)  {CATEGORY_HELP[cat][:60]}")
    print()
    print(
        f"structure hit (exact+IK1): "
        f"{sum(1 for r in rows if r['is_structure_hit_IK1_or_exact'])}/{n}"
    )
    print(
        f"near-or-better:            "
        f"{sum(1 for r in rows if r['is_near_miss_or_better'])}/{n}"
    )
    print(f"wrote {OUT_XLSX}")
    print(f"wrote {OUT_XLSX_SEALED}")
    print(f"wrote {OUT_CSV}")
    print(f"wrote {OUT_SUMMARY_CSV}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

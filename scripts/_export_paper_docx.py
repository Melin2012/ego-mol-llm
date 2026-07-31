#!/usr/bin/env python3
"""Export paper METHODS / RESULTS / HANDOFF / email draft to Word .docx."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "paper"


def set_styles(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for i in range(1, 4):
        h = doc.styles[f"Heading {i}"]
        h.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        h.font.name = "Calibri"


def add_title(doc: Document, text: str, subtitle: str | None = None) -> None:
    p = doc.add_heading(text, level=0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if subtitle:
        s = doc.add_paragraph(subtitle)
        s.runs[0].italic = True
        s.runs[0].font.size = Pt(10)
        s.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def add_bullets(doc: Document, items: list[str]) -> None:
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for it in items:
        doc.add_paragraph(it, style="List Number")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = str(val)
    doc.add_paragraph()


def build_methods() -> Path:
    doc = Document()
    set_styles(doc)
    add_title(
        doc,
        "Methods",
        "Network-assisted MS/MS structure annotation with free-form LLMs (ego-mol-llm). "
        "Manuscript-ready draft — July 2026.",
    )

    doc.add_heading("Overview", level=1)
    add_para(
        doc,
        "We developed a network-assisted MS/MS structure assignment workflow in which an "
        "unlabeled precursor is annotated using its ego neighborhood in a spectral molecular "
        "network, optional spectral library reverse search, offline fragment explanation "
        "(rule-based and/or CFM-ID), and a frontier chemistry LLM as a free-form reasoner over "
        "that evidence. The product design treats the network and spectrum as primary evidence "
        "and the LLM as the structure-calling step—not pure de novo elucidation without network "
        "context, and not closed-set ranking alone.",
    )
    add_para(
        doc,
        "Complementary ablations include deterministic product rankers (mass-gated neighbor/"
        "library voting), SIRIUS/CSI:FingerID, and CFM-ID network fragment transfer, used either "
        "as prompt blocks or as ranking signals.",
    )

    doc.add_heading("Molecular network and spectrum inputs", level=1)
    add_para(
        doc,
        "GraphML. Spectral networks were provided as GraphML (GNPS / HNSW-style). Nodes carried "
        "precursor mass (PEPMASS), optional library name and SMILES, community labels, and "
        "direct-neighbor flags. Edges stored MS/MS cosine similarity and absolute precursor "
        "delta m/z. For MassSpecGym HNSW egos, the query was encoded as a fixed seed node "
        "(id=9999999) with PEPMASS equal to the query precursor; library nodes retained names/SMILES.",
    )
    add_para(
        doc,
        "MGF. Seed and subgraph MS/MS spectra were supplied as MGF. Subgraph spectra were keyed "
        "by NETWORK_NODE_ID matching GraphML node ids. Seed spectra for blind packs used precursor "
        "m/z, ion mode, and peak lists only (structure metadata removed).",
    )
    add_para(
        doc,
        "Library reverse search (optional). When enabled, seed spectra were reverse-searched "
        "against NIST2023 MS/MS (LEVEL2 index) with precursor mass filtering and match-score "
        "ranking. Product policy treats seed NIST as optional verification (high score + mass-OK "
        "may support a structure; low scores are ignored). Neighbor-only NIST can be used for "
        "ablations without seed reverse search.",
    )

    doc.add_heading("Blind ego-context construction", level=1)
    add_numbered(
        doc,
        [
            "Load GraphML and resolve the seed (explicit seed id 9999999 for MSG HNSW; do not rely on hub heuristics).",
            "Build a one-hop (and optional two-hop) ego with seed name/SMILES hidden in blind mode.",
            "Rank neighbors by mass-aware evidence score (edge cosine, near-isobar/multimer residuals, annotation richness; default top N = 25–50).",
            "Attach spectral context: seed peaks, neighbor peaks, MS/MS cosine to seed, offline MS/MS explanation.",
        ],
    )
    add_para(
        doc,
        "Mass / adduct gate. Candidate structures must fit observed precursor m/z under even-electron "
        "monomer and multimer adducts, including multi-water losses when needed. Default monomer mass "
        "tolerance was 0.05 Da unless stated. Multimer residual gate default 0.10 Da.",
    )

    doc.add_heading("Free-form LLM structure assignment", level=1)
    add_para(
        doc,
        "Frontier models received a fixed system prompt specifying network-assisted annotation; "
        "mass-first rules; multimer awareness; use of MS/MS and dual cosine; optional NIST as soft "
        "evidence; required JSON schema (smiles, name, formula, adduct, confidence, rationale, alternatives).",
    )
    add_para(
        doc,
        "Strict free-form protocol required MASS → NETWORK → MS/MS → CHEMISTRY → DECISION and banned "
        "pack-wide SMILES indices, mass-transfer batch scripts, and external registry lookup. "
        "Strict-blind packs redacted database accessions from model-facing files; sealed truth retained "
        "full linkage offline.",
    )
    add_para(
        doc,
        "Neighbor-only NIST evaluation arm. Neighbor MS/MS peaks were reverse-searched against NIST2023; "
        "hits were injected as a prompt block. The seed was never reverse-searched. Models were instructed "
        "to promote a neighbor NIST structure only if mass-consistent with the seed precursor.",
    )
    add_para(
        doc,
        "Free-form validity. Valid free-form arms require sample-by-sample model consumption of each prompt. "
        "Bulk writers that mass-fit SMILES from pack prompts in minutes (or that self-report m/z clustering) "
        "were archived as ranker-style ablations and excluded from free-form accuracy claims.",
    )

    doc.add_heading("Blind evaluation packs", level=1)
    doc.add_heading("ASTRAL / ego holdout packs (n = 40)", level=2)
    add_para(
        doc,
        "Randomized holdouts (e.g. holdout-40c/40d) were drawn from ASTRAL C18 networks. Each pack "
        "included GraphML, blind seed MGF, subgraph MGF, jobs/prompts, and a sealed truth index. "
        "Evidence arms included pure network, MGF-only ± NIST, SIRIUS-first, CFM-first, and product "
        "neighbor+MGF ranker.",
    )

    doc.add_heading("MassSpecGym HNSW subset (n = 285)", level=2)
    add_para(
        doc,
        "Spectra with precomputed HNSW ego GraphML and subgraph MGF were linked by full-file index N "
        "(HNSW_spectrum_N ↔ position N in MassSpecGym.mgf / official spectrum index). Overlap with the "
        "full official test fold was partial (285 of 17,556 test spectra in the early HNSW index range). "
        "A strict v2 rebuild removed MassSpecGym accessions from model-facing prompts.",
    )

    doc.add_heading("MassSpecGym HNSW full block (n = 11,540)", level=2)
    add_para(
        doc,
        "A second HNSW dump covered indices 219564–231103 (11,540 contiguous file-order positions). "
        "Seed PEPMASS checks matched the official index at 100%. Under official MSG fold labels "
        "(joined after index N; not present in GraphML), this block comprises approximately "
        "11,386 train / 90 val / 64 test spectra—not the full official test fold (17,556). "
        "The block is chemically redundant (~1,780 unique IK1; ~6.5 spectra per molecule). "
        "Recommended free-form API set: official fold=test only, no IK1 dedupe (n=64; "
        "scripts/subset_msg_hnsw_handout_by_fold.py --fold test).",
    )

    doc.add_heading("Scoring", level=1)
    add_bullets(
        doc,
        [
            "IK1: InChIKey first block equality (connectivity).",
            "Exact: canonical SMILES equality.",
            "Formula: RDKit molecular formula equality.",
            "Tanimoto: Morgan fingerprints, radius 2, 2048 bits (thresholds 0.7 and 0.85).",
        ],
    )
    add_para(
        doc,
        "When sealed InChIKeys disagreed with RDKit(true_smiles), scoring used RDKit-derived keys from true SMILES. "
        "Unique-molecule metrics are recommended when spectrum-level counts over-represent frequent structures.",
    )

    doc.add_heading("Case study: library meta-tyrosine mislabeled as tyrosine", level=1)
    add_para(
        doc,
        "Holdout spectrum AROMEC18COLGATE000057 was named “Tyrosine” with formula C9H11NO3, but sealed SMILES "
        "corresponds to meta-tyrosine (3-OH). Models often proposed para-tyrosine (standard L-Tyr), scored as "
        "IK1 miss with formula match—regioisomer disagreement with a misnamed library structure, not a random miss. "
        "This motivates scoring against structure (InChIKey/SMILES), not free-text names alone.",
    )
    add_table(
        doc,
        ["Structure", "OH position", "Role"],
        [
            ["meta-tyrosine (sealed SMILES)", "meta (3-)", "Library / sealed truth"],
            ["L-tyrosine (common name)", "para (4-)", "Typical model prediction"],
        ],
    )

    doc.add_heading("Baselines and ablations", level=1)
    add_numbered(
        doc,
        [
            "Deterministic product ranker: mass-OK neighbor SMILES; seed NIST only if score ≥ 0.85 and mass-OK.",
            "MGF-only ± NIST (no GraphML).",
            "SIRIUS/CSI-first and CFM-first ranking without free-form LLM.",
            "Free-form multi-model comparison on frozen prompts (Opus, Grok, local ChemDFM/Qwen).",
        ],
    )

    doc.add_heading("Limitations", level=1)
    add_para(
        doc,
        "Network LLM annotation inherits library annotation errors and spectral isobar confusion. "
        "MassSpecGym HNSW results apply only where ego GraphML exists; partial coverage must not be reported "
        "as full official-test performance. Free-form API cost limits complete multi-model runs on large packs. "
        "Expert review remains required for publication-grade IDs.",
    )

    doc.add_heading("Software and data availability", level=1)
    add_para(
        doc,
        "Implementation: ego-mol-llm (Apache-2.0), Python 3.10+, RDKit; optional SIRIUS CLI and CFM-ID Docker. "
        "Repository: https://github.com/Melin2012/ego-mol-llm. Blind packs use prompts/, jobs/, predictions_<model>/, "
        "and sealed truth_index.csv offline.",
    )

    path = OUT / "METHODS.docx"
    doc.save(path)
    return path


def build_results() -> Path:
    doc = Document()
    set_styles(doc)
    add_title(
        doc,
        "Results Board (Frozen Evaluation Snapshot)",
        "July 2026. “Valid free-form” = sample-by-sample LLM reasoning without pack-wide mass-fit automation.",
    )

    doc.add_heading("A. Product algorithm (final)", level=1)
    add_numbered(
        doc,
        [
            "Inputs: GraphML ego (seed 9999999 for MSG HNSW), subgraph MGF (NETWORK_NODE_ID), seed MGF.",
            "Blind ego: hide seed name/SMILES; rank neighbors; offline MS/MS explanation.",
            "Optional neighbor-only NIST (seed NIST OFF).",
            "Free-form LLM: MASS→NETWORK→MS/MS→CHEMISTRY→DECISION; one JSON SMILES; no pack-wide index.",
            "Optional product ranker / hybrid (seed NIST only if score ≥ 0.85).",
            "Score offline: IK1, exact SMILES, formula, Morgan-2 T ≥ 0.7 / 0.85.",
        ],
    )
    add_para(
        doc,
        "Link rule (MSG HNSW): HNSW_spectrum_N ↔ full MassSpecGym.mgf file-order index N. "
        "Official fold labels come from the MSG index after that join—not from GraphML/MGF.",
        bold=False,
    )

    doc.add_heading("B. Holdout packs (ASTRAL C18, n = 40)", level=1)
    add_para(
        doc,
        "Product neighbor+MGF ranker ~60% IK1; SIRIUS-first ~55–57.5%; network ~40–43%; "
        "MGF+NIST inflated when seed library is unrestricted. Tyr meta/para library case in Methods.",
    )

    doc.add_heading("C. MassSpecGym HNSW subset — strict blind 285", level=1)
    add_table(
        doc,
        ["Item", "Value"],
        [
            ["n spectra", "285"],
            ["Unique true IK1", "18 (high replicate rate)"],
            ["Coverage of official MSG test", "285 / 17,556 (old HNSW index range only)"],
            ["Model-facing pack", "MSG_HNSW_blind285v2_strict_ego_msms"],
            ["Sealed truth", "MSG_HNSW_blind285v2_strict_ego_msms_SEALED_truth"],
        ],
    )

    doc.add_heading("C1. Free-form / local models", level=2)
    add_table(
        doc,
        ["Arm", "n scored", "IK1", "Exact", "Notes"],
        [
            ["Grok free-form, no neighbor NIST (strict v2)", "285", "102 (35.8%)", "90", "Credible free-form; empty SMILES 128; unique-mol any-hit 13/18"],
            ["Claude Opus free-form + neighbor NIST", "64 / 285", "55 (85.9%)", "52", "Partial pilot; ~8M tokens/64; killed — unsustainable"],
            ["Qwen3:14b free-form + neighbor NIST", "87 / 285", "12 (13.8%)", "12", "Partial; max-runtime stop; 12/40 assigned = 30% IK1"],
            ["ChemDFM-R pure model", "285", "3 (1.1%)", "2", "Mass-illegal hallucinations"],
        ],
    )

    doc.add_heading("C2. Invalid / ablation (do NOT cite as free-form)", level=2)
    add_table(
        doc,
        ["Arm", "IK1", "Why invalid"],
        [
            ["Grok ranker-fast nist neigh", "195/285 (68.4%)", "Helper scripts / bulk mass-fit"],
            ["Grok online m/z-cluster zip", "187/285 (65.6%)", "Self-admitted bulk m/z clustering"],
            ["Grok clone of no-NIST freeform", "100/285 (35.1%)", "~10 s relabel of prior free-form"],
            ["v1 free-form 249–252/285", "—", "Leaky pack; superseded by strict v2"],
        ],
    )

    doc.add_heading("C3. Outcome buckets (online Grok m/z-cluster, ranker upper bound)", level=2)
    add_para(doc, "Exact 175 · Similar 16 · Formula-only 10 · True miss 20 · Empty 64 (n=285).")

    doc.add_heading("D. MassSpecGym HNSW full block 11,540 (new dump)", level=1)
    add_table(
        doc,
        ["Item", "Value"],
        [
            ["Ego nets", "11,540 (HNSW_spectrum_219564…231103)"],
            ["Official MSG test fold size", "17,556"],
            ["Match full test?", "No"],
            ["Official folds inside this block", "train 11,386 / val 90 / test 64"],
            ["Unique IK1", "1,780 (~6.5× spectra per molecule)"],
            ["Top 100 mols cover", "66.5% of spectra"],
            ["Seed PEPMASS vs MSG precursor", "11,540/11,540 (index-N link OK)"],
        ],
    )

    doc.add_heading("D0. Recommended free-form: test fold only (n=64, no dedupe)", level=2)
    add_para(
        doc,
        "All official fold=test spectra from the HNSW block — original test count preserved "
        "(not IK1-deduped). Script: scripts/subset_msg_hnsw_handout_by_fold.py --fold test.",
    )
    add_table(
        doc,
        ["Item", "Value"],
        [
            ["n spectra", "64 (= parent test count)"],
            ["Unique IK1 inside", "53"],
            ["Dedupe?", "No"],
            ["Handout", "Desktop\\MSG_HNSW_test64_nist_neigh_handout"],
            ["Sealed", "Desktop\\MSG_HNSW_test64_nist_neigh_SEALED_truth"],
            ["Grok zip", "Desktop\\MSG_HNSW_test64_nist_neigh_GROK_handout.zip"],
        ],
    )
    add_para(
        doc,
        "Still not full official MSG test (17,556). Optional larger packs: IK1-deduped 1,780; full 11,540.",
    )

    doc.add_heading("D1. Grok return on full 11,540 (INVALID free-form)", level=2)
    add_table(
        doc,
        ["Item", "Value"],
        [
            ["Wall time", "~1.6 min for 11,540 → bulk automation"],
            ["IK1 all", "1,105/11,540 (9.6%)"],
            ["Empty SMILES", "8,349 (72%)"],
            ["IK1 official test subset", "21/64 (32.8%)"],
            ["Use", "Ranker-style ablation only — not free-form"],
        ],
    )

    doc.add_heading("D2. Frontier API cost order (valid free-form)", level=2)
    add_para(doc, "Rough planning (input ~8k tok/sample; solid free-form ~4k out):")
    add_table(
        doc,
        ["Model (API)", "Full 11,540", "Deduped 1,780"],
        [
            ["Claude Opus 5", "~$1.5k–$5k+", "~$0.2k–$0.8k+"],
            ["Claude Sonnet 5 (intro)", "~$0.6k–$2k", "~$90–$300"],
            ["Grok 4.5 / 4.3", "~$0.25k–$1.2k", "~$40–$180"],
            ["DeepSeek V4 Flash", "~$50–$400", "~$8–$60"],
        ],
    )
    add_para(
        doc,
        "Recommendation: free-form on deduped 1,780 (or ids_test first); meter a 20-sample pilot; "
        "prefer Sonnet or Grok API.",
    )

    doc.add_heading("E. Student TODO (benchmarks)", level=1)
    add_numbered(
        doc,
        [
            "Finish valid free-form on frozen prompts (285 strict and/or deduped 1,780).",
            "Deterministic product ranker on same packs (neighbor mass-OK ± neighbor NIST).",
            "Optional SIRIUS/CFM ablations already scripted.",
            "Tables: spectrum-level (full 11,540) + unique-IK1 (dedup pack); never mix invalid bulk with free-form.",
            "When new full MSG HNSW test coverage arrives, re-link and re-run under same protocol.",
        ],
    )

    path = OUT / "RESULTS_BOARD.docx"
    doc.save(path)
    return path


def build_handoff() -> Path:
    doc = Document()
    set_styles(doc)
    add_title(
        doc,
        "Handoff: ego-mol-llm evaluation → Alexander / student",
        "From Alexey (product + packs + protocol freeze). Repo: https://github.com/Melin2012/ego-mol-llm",
    )

    doc.add_heading("1. Algorithm (product path — frozen)", level=1)
    add_para(
        doc,
        "GraphML ego + seed MGF + subgraph MGF → hide seed structure → neighbor ranking + offline MS/MS "
        "explain → optional NEIGHBOR-ONLY NIST (seed NIST OFF) → free-form LLM "
        "(MASS→NETWORK→MS/MS→CHEMISTRY→DECISION) OR mass-OK ranker → JSON SMILES → score vs SEALED truth "
        "(IK1 / exact / formula / T≥0.7).",
    )
    add_table(
        doc,
        ["Rule", "Detail"],
        [
            ["MSG HNSW seed node", "Always 9999999 — do not use find_seed() hub heuristic"],
            ["Link GraphML → MSG", "HNSW_spectrum_N ↔ file-order index N in full MassSpecGym.mgf / spectrum index"],
            ["Fold labels", "From official MSG index after join — not inside GraphML/MGF"],
            ["Neighbor NIST", "Reverse-search neighbor peaks only; inject NIST block into prompt"],
            ["Free-form validity", "One sample at a time; ban pack-wide SMILES index / m/z-cluster bulk"],
            ["Subgraph MGF", "Neighbors only; seed MS/MS from full MSG (or blind seed MGF)"],
        ],
    )

    doc.add_heading("2. Code map (repo)", level=1)
    add_table(
        doc,
        ["Area", "Path"],
        [
            ["Core library", "src/ego_mol_llm/"],
            ["Strict 285 pack builder", "scripts/build_msg_hnsw_blind285_strict.py"],
            ["Neighbor-NIST refresh", "scripts/refresh_pack_nist_neighbors_only.py"],
            ["Full 11,540 handout builder", "scripts/build_msg_hnsw_full_new_nist_handout.py"],
            ["Fold subset (test64)", "scripts/subset_msg_hnsw_handout_by_fold.py"],
            ["IK1 dedupe handout (optional)", "scripts/dedupe_msg_hnsw_handout_by_ik1.py"],
            ["Free-form runners", "scripts/run_jobs_prompt_ollama.py, run_blind_pack_ollama.py"],
            ["Product ranker pattern", "scripts/run_product_neighbor_mgf_40c.py"],
            ["MSG split", "scripts/split_massspecgym_mgf.py"],
            ["Methods / board (Word + MD)", "paper/METHODS.docx, RESULTS_BOARD.docx, …"],
        ],
    )
    add_para(doc, "Install: pip install -e \".[api]\" (and RDKit for scoring).")

    doc.add_heading("3. Data locations (Alexey’s machine)", level=1)
    add_table(
        doc,
        ["Dataset", "Path"],
        [
            ["Strict 285 pack", "Desktop\\MSG_HNSW_blind285v2_strict_ego_msms"],
            ["Strict 285 sealed", "Desktop\\MSG_HNSW_blind285v2_strict_ego_msms_SEALED_truth"],
            ["Full 11,540 handout", "Desktop\\MSG_HNSW_full11540_nist_neigh_handout"],
            ["Full 11,540 sealed", "Desktop\\MSG_HNSW_full11540_nist_neigh_SEALED_truth"],
            ["Grok prompts zip (full)", "Desktop\\MSG_HNSW_full11540_nist_neigh_GROK_handout.zip"],
            ["Test-only handout (n=64, no dedupe)", "Desktop\\MSG_HNSW_test64_nist_neigh_handout"],
            ["Test-only sealed", "Desktop\\MSG_HNSW_test64_nist_neigh_SEALED_truth"],
            ["Grok zip (test64)", "Desktop\\MSG_HNSW_test64_nist_neigh_GROK_handout.zip"],
            ["Fold subset script", "scripts/subset_msg_hnsw_handout_by_fold.py"],
            ["IK1-deduped optional (n=1780)", "Desktop\\MSG_HNSW_dedup_ik1_nist_neigh_handout"],
            ["Dedupe script", "scripts/dedupe_msg_hnsw_handout_by_ik1.py"],
            ["Raw GraphML/MGF", "HNSW_Large_Files\\graphmls_new\\graphmls (+ subgraph_mgfs_new)"],
            ["MSG split / index", "Downloads\\MassSpecGym_split\\"],
            ["QC + redundancy", "HNSW_Large_Files\\MSG_FULL_NEW_ANALYSIS\\"],
            ["NIST index", "Codes\\ego-mol-llm\\libraries\\LEVEL2_NIST2023MSMS_20240408.index.pkl"],
        ],
    )
    add_para(doc, "Do not ship sealed truth to external model operators.", bold=True)

    doc.add_heading("4. Frozen numbers (summary)", level=1)
    add_bullets(
        doc,
        [
            "Valid free-form baseline (285 strict, no neigh NIST): Grok IK1 102/285 (35.8%).",
            "Claude free-form + neigh NIST: 64/285 only, IK1 85.9% interim — token cost killed full run.",
            "Qwen3:14b + neigh NIST: 87/285, IK1 13.8% (partial).",
            "Invalid Grok bulk on 285 and 11,540: not free-form (seconds–minutes for full packs).",
            "11,540 block ≠ official MSG test (17,556); official folds inside block: 11386/90/64.",
            "Redundancy: 11,540 spectra → 1780 unique IK1 (~6.5×).",
            "Preferred free-form handout: test-only n=64 (original test count, no IK1 dedupe).",
        ],
    )

    doc.add_heading("5. Student checklist", level=1)
    doc.add_heading("Required", level=2)
    add_bullets(
        doc,
        [
            "Reproduce scoring on sealed truth (IK1 / exact / formula / T≥0.7).",
            "Run deterministic product ranker on frozen 285 (and optionally 11,540).",
            "Complete one valid free-form arm on 285 with neighbor-NIST prompts (API frontier with output caps) or document cost stop.",
            "Report spectrum-level and unique-molecule metrics.",
            "Separate tables: free-form vs ranker vs invalid bulk.",
        ],
    )
    doc.add_heading("Strongly recommended", level=2)
    add_bullets(
        doc,
        [
            "Free-form on deduped pack (n=1780; train 1652 / val 75 / test 53) or start with ids_test.txt (n=53).",
            "Official test-only rows in 11,540 block (n=64) as secondary table.",
            "When new full MSG HNSW test coverage arrives: re-link with same seed/N protocol.",
        ],
    )
    doc.add_heading("Do not", level=2)
    add_bullets(
        doc,
        [
            "Cite 5-minute / 1.6-minute Grok dumps as free-form.",
            "Use find_seed() instead of 9999999 on MSG HNSW.",
            "Put sealed SMILES in model-facing zips.",
        ],
    )

    doc.add_heading("6. Ownership", level=1)
    add_table(
        doc,
        ["Role", "Person"],
        [
            ["Product, packs, protocol, invalid-run audit", "Alexey"],
            ["Benchmarks, student execution, final tables", "Alexander → student"],
            ["Code repo", "Melin2012/ego-mol-llm"],
        ],
    )

    path = OUT / "HANDOFF_ALEXANDER.docx"
    doc.save(path)
    return path


def build_email() -> Path:
    doc = Document()
    set_styles(doc)
    add_title(doc, "Draft email to Alexander", "Forward to student for finishing benchmarks.")

    add_para(doc, "Subject:", bold=True)
    add_para(
        doc,
        "ego-mol-llm handoff — frozen protocol, packs, results board, student benchmarks TODO",
    )
    doc.add_paragraph()
    add_para(doc, "Hi Alexander,")
    add_para(
        doc,
        "Please pass this package to your student to finish the benchmarks for the network + free-form "
        "LLM MS/MS structure paper. Product path, blind packs, and evaluation protocol are frozen on my "
        "side; remaining work is clean scoring tables and completing free-form / ranker arms under the same rules.",
    )
    add_para(doc, "Repo:", bold=True)
    add_para(doc, "https://github.com/Melin2012/ego-mol-llm")
    add_para(
        doc,
        "Key docs (also as Word files in paper/): METHODS.docx, RESULTS_BOARD.docx, HANDOFF_ALEXANDER.docx.",
    )
    add_para(doc, "Algorithm (one paragraph):", bold=True)
    add_para(
        doc,
        "Unknown seed in a spectral ego-network (GraphML + MGF) → hide seed structure → rank neighbors → "
        "optional neighbor-only NIST (seed never library-searched) → free-form LLM reasons "
        "MASS→NETWORK→MS/MS→CHEMISTRY→DECISION and outputs one SMILES JSON — scored offline vs sealed truth "
        "(IK1 / exact SMILES / formula / Tanimoto). Deterministic mass-OK ranker is a separate ablation. "
        "For MassSpecGym HNSW, seed node is always 9999999; link to MSG is by file-order index N. Official "
        "train/val/test folds come from the MSG index after that join (not from GraphML).",
    )
    add_para(doc, "What is ready:", bold=True)
    add_numbered(
        doc,
        [
            "Strict blind 285 pack (HNSW-available MSG subset) + sealed truth + neighbor-NIST prompts.",
            "Full 11,540 HNSW block handout with neighbor NIST precomputed + sealed truth + Grok prompt zip (~89 MB).",
            "QC: index-N link 100% PEPMASS match; IK1-deduped free-form set n=1780 (train/val/test 1652/75/53).",
            "Local runners (Ollama) + product ranker scripts in repo.",
        ],
    )
    add_para(doc, "Important caveats (please stress to the student):", bold=True)
    add_numbered(
        doc,
        [
            "11,540 ≠ official MSG test (17,556). Under official MSG folds this block is mostly train (11,386 / 90 / 64). We will re-run when fuller HNSW coverage of the true test fold is available.",
            "Several fast Grok “full pack” returns are invalid free-form (bulk mass-fit / m/z cluster in minutes). Do not put them in free-form tables. Credible free-form baseline on 285: Grok ~36% IK1 without neighbor NIST. Claude pilot ~86% on 64/285 then stopped (~8M tokens). Qwen local ~14% on 87/285 partial.",
            "Free-form on all 11,540 via frontier API is the intended completion path (Sonnet/Grok-class cost order hundreds–low thousands USD if thinking is controlled; Opus can be much higher). Deduping to 1 spectrum per IK1 (~1.8k) cuts cost ~6×.",
        ],
    )
    add_para(doc, "Student asks:", bold=True)
    add_numbered(
        doc,
        [
            "Finish ranker + free-form tables on frozen packs (spectrum + unique-molecule metrics).",
            "Keep free-form vs ranker vs invalid bulk in separate columns.",
            "Coordinate API budget with you; I can help wire the batch runner if needed.",
        ],
    )
    add_para(
        doc,
        "Happy to jump on a short call with you + the student to walk through folder layout and scoring.",
    )
    add_para(doc, "Best,")
    add_para(doc, "Alexey")

    path = OUT / "EMAIL_DRAFT_ALEXANDER.docx"
    doc.save(path)
    return path


def main() -> None:
    OUT.mkdir(exist_ok=True)
    paths = [build_methods(), build_results(), build_handoff(), build_email()]
    for p in paths:
        print("wrote", p)


if __name__ == "__main__":
    main()
